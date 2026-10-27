from pptx import Presentation
import sys

prs = Presentation(r'C:\Users\navin\Downloads\WindTunnel_Review2.pptx')

with open('all_slides_summary.txt', 'w', encoding='utf-8') as out:
    for idx, slide in enumerate(prs.slides, 1):
        out.write(f"\n==================== SLIDE {idx} ====================\n")
        for shape in slide.shapes:
            stype = shape.shape_type
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    if p.text.strip():
                        out.write(f"  [TEXT: {shape.name}] {p.text.strip()}\n")
            elif shape.has_table:
                for row in shape.table.rows:
                    row_str = ' | '.join([c.text.strip().replace('\n', ' ') for c in row.cells if c.text.strip()])
                    if row_str:
                        out.write(f"  [TABLE: {shape.name}] {row_str}\n")
            elif stype == 3: # chart
                title = shape.chart.chart_title.text_frame.text if shape.chart.has_title else "No Title"
                out.write(f"  [CHART: {shape.name}] Title: {title}\n")
            elif stype == 13: # picture
                out.write(f"  [PICTURE: {shape.name}]\n")
            else:
                out.write(f"  [SHAPE ({stype}): {shape.name}]\n")

print("Saved all_slides_summary.txt")
