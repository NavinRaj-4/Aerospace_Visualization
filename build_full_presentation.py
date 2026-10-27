import os
import shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def build_presentation():
    ref_path = r'C:\Users\navin\Downloads\WindTunnel_Review2.pptx'
    prs = Presentation(ref_path)
    
    # Colors
    TEAL = RGBColor(0, 168, 150)       # #00A896
    CYAN = RGBColor(2, 128, 144)       # #028090
    DARK_TEXT = RGBColor(27, 42, 46)   # #1B2A2E
    MUTED_TEXT = RGBColor(122, 143, 140)# #7A8F8C
    WHITE = RGBColor(255, 255, 255)
    
    # Helper to set text in a shape
    def set_shape_text(shape, lines, font_size=None, font_color=None, bold=False):
        tf = shape.text_frame
        tf.clear()
        for idx, line in enumerate(lines):
            p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            p.text = line
            if font_size:
                p.font.size = font_size
            if font_color:
                p.font.color.rgb = font_color
            if bold is not None:
                p.font.bold = bold

    # Slide 1: Title Slide (Preserve institution, guide, presenter, refine title if needed)
    # Already perfect.
    
    # Slide 2: Agenda
    slide2 = prs.slides[1]
    for s in slide2.shapes:
        if s.name == "Text 3":
            agenda_items = [
                "1. Project Scope & Aerodynamic Requirements (Functional & Non-Functional)",
                "2. Detailed Multi-Tier System Architecture & Telemetry Data Flow",
                "3. Aerodynamic System Use Case Diagram & Actor Interactions",
                "4. ER Diagram & Aerospace Entity Data Models (Models, Runs, Pressure Taps)",
                "5. Modular Subsystem Breakdown & Key Telemetry Functions",
                "6. Modern Full-Stack Technology Stack & Toolchain (React, Vite, Node, Canvas)",
                "7. UI Wireframe & Telemetry HUD Design System",
                "8. Frontend Engineering (React Components, 3D Canvas Flow Engine, Audio Synth)",
                "9. Aerodynamic Data Visualization (Live CL, CD, Cp, L/D Output)",
                "10. Backend Engineering & Dynamic Aerodynamic Calculation Engine",
                "11. Database / In-Memory Telemetry Integration & Multi-Regime Datasets",
                "12. RESTful API Telemetry Endpoints & Request-Response Pipeline",
                "13. Module-wise Implementation Progress Status (85% – 95%)",
                "14. Live Interface Previews & Interactive Simulator Screenshots",
                "15. Technical Challenges Faced & Physics Solutions Adopted",
                "16. Future Scope, Cloud Persistence & DAQ Hardware Integration",
                "17. References & Academic Aerodynamic Citations"
            ]
            set_shape_text(s, agenda_items, font_size=Pt(11.5), font_color=WHITE)

    # Slide 3: Requirements
    slide3 = prs.slides[2]
    for s in slide3.shapes:
        if s.name == "Text 7":
            func_items = [
                "• Real-time aerodynamic telemetry computation (CL, CD, CM, Cp, L/D, q_inf)",
                "• Multi-model airfoil management (NACA 0012, NASA SC(2)-0714, 65° Delta Wing)",
                "• Multi-regime wind tunnel facility configuration (Subsonic, Transonic, Supersonic)",
                "• Interactive 3D Wind Tunnel Canvas with flow particles, streamlines & stall wake",
                "• 5-Fold Aerodynamic Analytics Suite (CL-α, CD-α, Drag Polar, Cp-x/c, L/D-α)",
                "• Automated Live AoA Pitch Sweep Engine (-4° to +20°) with Web Audio telemetry",
                "• Dataset Inspector for CSV/JSON ingestion, validation, filtering, and export",
                "• High-performance RESTful telemetry API with in-memory fast-query store"
            ]
            set_shape_text(s, func_items, font_size=Pt(10.0), font_color=DARK_TEXT)
        elif s.name == "Text 9":
            non_func_items = [
                "• Real-time responsiveness (< 5ms API latency, 60 FPS Canvas rendering)",
                "• High-accuracy aerodynamic formulations (Prandtl-Glauert compressibility)",
                "• Responsive dark aerospace HUD theme (Tailwind CSS, Lucide Icons)",
                "• Decoupled stateless REST architecture with modular components",
                "• Robust dataset ingestion with automated unit normalization and validation",
                "• Acoustic audio synthesis using Web Audio API harmonic oscillators"
            ]
            set_shape_text(s, non_func_items, font_size=Pt(10.0), font_color=DARK_TEXT)

    # Slide 4: System Architecture
    slide4 = prs.slides[3]
    for s in slide4.shapes:
        if s.name == "Text 6":
            set_shape_text(s, ["React 18 + Vite 6", "HTML5 Canvas Flow Engine", "Chart.js / react-chartjs-2", "Web Audio Synthesizer"], font_size=Pt(9.5), font_color=WHITE)
        elif s.name == "Text 10":
            set_shape_text(s, ["Node.js + Express.js", "Aerodynamic Physics Engine", "Live AoA Sweep Generator", "CSV Dataset Ingestion Parser"], font_size=Pt(9.5), font_color=WHITE)
        elif s.name == "Text 14":
            set_shape_text(s, ["Aerospace Models & Polars", "Tunnel Facilities Telemetry", "Test Run In-Memory Store", "13-Station Cp Profiles"], font_size=Pt(9.5), font_color=WHITE)

    # Slide 7: Module Description Table
    slide7 = prs.slides[6]
    for s in slide7.shapes:
        if s.has_table:
            table = s.table
            modules_data = [
                ("3D Flow Simulator", "Interactive Canvas aerodynamic particle simulator", "Real-time streamlines, stall vortex shedding, smoke rakes"),
                ("Aerodynamic Analytics", "5-fold scientific performance visualization suite", "CL-α lift curve, Drag Polar, Cp distribution, L/D efficiency"),
                ("Live AoA Sweep Engine", "Automated motorized angle of attack sweeps", "Continuous angle traversing, streaming telemetry, audio synth"),
                ("Facility Mission Control", "Multi-regime wind tunnel facility monitoring", "RPM, dynamic pressure (q_inf), cryogenic cooling status"),
                ("Dataset Inspector", "Flight test run data exploration & ingestion", "Batch CSV upload, multi-parameter filtering, JSON/CSV export"),
                ("Telemetry REST API", "High-speed backend aerodynamic calculation server", "Real-time coefficient solver, sweep generator, health metrics")
            ]
            for r_idx, (m, desc, fn) in enumerate(modules_data, 1):
                if r_idx < len(table.rows):
                    table.rows[r_idx].cells[0].text = m
                    table.rows[r_idx].cells[1].text = desc
                    table.rows[r_idx].cells[2].text = fn
                    for c in table.rows[r_idx].cells:
                        for p in c.text_frame.paragraphs:
                            p.font.size = Pt(9.5)
                            p.font.color.rgb = DARK_TEXT

    # Slide 8: Technology Stack
    slide8 = prs.slides[7]
    for s in slide8.shapes:
        if s.name == "Text 5":
            s.text_frame.text = "React 18 & Vite"
            s.text_frame.paragraphs[0].font.size = Pt(13)
            s.text_frame.paragraphs[0].font.bold = True
            s.text_frame.paragraphs[0].font.color.rgb = CYAN
        elif s.name == "Text 6":
            set_shape_text(s, ["High-speed frontend framework with instant HMR and component modularity."], font_size=Pt(10), font_color=DARK_TEXT)
        elif s.name == "Text 8":
            s.text_frame.text = "HTML5 Canvas"
            s.text_frame.paragraphs[0].font.size = Pt(13)
            s.text_frame.paragraphs[0].font.bold = True
            s.text_frame.paragraphs[0].font.color.rgb = CYAN
        elif s.name == "Text 9":
            set_shape_text(s, ["Real-time 60 FPS aerodynamic vector field & particle streamline engine."], font_size=Pt(10), font_color=DARK_TEXT)
        elif s.name == "Text 11":
            s.text_frame.text = "Chart.js 4"
            s.text_frame.paragraphs[0].font.size = Pt(13)
            s.text_frame.paragraphs[0].font.bold = True
            s.text_frame.paragraphs[0].font.color.rgb = CYAN
        elif s.name == "Text 12":
            set_shape_text(s, ["Interactive scientific charts rendering lift, drag, pitching moment & Cp loops."], font_size=Pt(10), font_color=DARK_TEXT)
        elif s.name == "Text 14":
            s.text_frame.text = "Node & Express"
            s.text_frame.paragraphs[0].font.size = Pt(13)
            s.text_frame.paragraphs[0].font.bold = True
            s.text_frame.paragraphs[0].font.color.rgb = CYAN
        elif s.name == "Text 15":
            set_shape_text(s, ["REST API telemetry server and dynamic aerodynamic calculation algorithms."], font_size=Pt(10), font_color=DARK_TEXT)
        elif s.name == "Text 17":
            tools_txt = [
                "• Tailwind CSS: Dark aerospace HUD theme with custom neon telemetry accents.",
                "• Web Audio API: Dynamic sound synthesis simulating aerodynamic airflow turbulence.",
                "• Lucide React Icons & Tooling: VS Code, Postman for API testing, Git & GitHub.",
                "• Python & NumPy: Aerodynamic numerical validation and curve modeling."
            ]
            set_shape_text(s, tools_txt, font_size=Pt(11.0), font_color=DARK_TEXT)

    # Slide 10: Frontend Development
    slide10 = prs.slides[9]
    for s in slide10.shapes:
        if s.name == "Text 4":
            front_items = [
                "• Built on React 18 and Vite 6 for ultra-fast compilation and zero-lag rendering",
                "• Interactive 3D Wind Tunnel Canvas with dynamic particle streamlines & stall wake shedding",
                "• Real-Time AoA Pitch Controller (-4° to +20°) dynamically recalculating telemetry in <1ms",
                "• Aerodynamic Analytics Suite featuring 5 synchronized Chart.js scientific polars",
                "• Live Motorized Sweep Controller with automated AoA traversing & Web Audio sound feedback",
                "• Dataset Inspector with multi-filter search, CSV file upload, and JSON telemetry export",
                "• Dark Aerospace Mission HUD styled with Tailwind CSS and responsive layout grids"
            ]
            set_shape_text(s, front_items, font_size=Pt(11.0), font_color=DARK_TEXT)
        elif s.name == "Text 6":
            set_shape_text(s, [
                "Platform Tabs & Views",
                "Mission Overview | 3D Wind Tunnel Simulator | Analytics Suite | Dataset Inspector | Live Tunnel Sweep"
            ], font_size=Pt(10.5), font_color=CYAN)
            s.text_frame.paragraphs[1].font.color.rgb = MUTED_TEXT

    # Slide 11: Replace slide 11 charts with real high-res scientific charts
    slide11 = prs.slides[10]
    # Remove existing chart shapes if possible or overlay
    # Let's inspect shapes in slide11
    shapes_to_remove = [s for s in slide11.shapes if s.name in ['Chart 0', 'Chart 1'] or s.shape_type == 3]
    for shp in shapes_to_remove:
        sp = shp._element
        sp.getparent().remove(sp)
    
    # Add 2 high-res chart images side-by-side
    slide11.shapes.add_picture('ppt_assets/chart_lift_curve.png', Inches(0.5), Inches(1.35), Inches(4.35), Inches(3.2))
    slide11.shapes.add_picture('ppt_assets/chart_drag_polar.png', Inches(5.15), Inches(1.35), Inches(4.35), Inches(3.2))
    for s in slide11.shapes:
        if s.name == "Text 4":
            set_shape_text(s, ["Live scientific charts calculated dynamically via backend telemetry engine (NACA 0012 & NASA SC(2)-0714)."], font_size=Pt(9.5), font_color=MUTED_TEXT)

    # Slide 12: Backend Development
    slide12 = prs.slides[11]
    for s in slide12.shapes:
        if s.name == "Text 4":
            back_items = [
                "• High-performance Node.js runtime and modular Express.js REST routing architecture",
                "• Dynamic Aerodynamic Telemetry Engine computing CL, CD, CM, and L/D on-the-fly",
                "• Compressibility Corrections implementing Prandtl-Glauert factor 1/sqrt(1 - M^2)",
                "• Supercritical Shock-Buffet & Post-Stall Separation modeling past critical angle",
                "• Continuous AoA Sweep Generator producing high-density telemetry data arrays",
                "• Multi-format CSV Dataset Ingestion with automatic unit normalization and error validation",
                "• High-speed in-memory state store delivering sub-millisecond query performance"
            ]
            set_shape_text(s, back_items, font_size=Pt(11.0), font_color=DARK_TEXT)
        elif s.name == "Text 6":
            set_shape_text(s, [
                "Backend Processing Flow",
                "REST Telemetry Endpoint → Physics Engine (calculateAerodynamics) → Unit Normalization → In-Memory Store → JSON Stream"
            ], font_size=Pt(10.5), font_color=CYAN)
            s.text_frame.paragraphs[1].font.color.rgb = MUTED_TEXT

    # Slide 13: Database & Datasets
    slide13 = prs.slides[12]
    for s in slide13.shapes:
        if s.name == "Text 4":
            db_items = [
                "• Multi-model aerospace registry storing geometric specifications and aerodynamic polars.",
                "• Multi-regime wind tunnel facility status tracking active RPM, dynamic pressure, and cooling.",
                "• In-memory fast-query store containing 156+ recorded test runs for instantaneous retrieval.",
                "• 13-station normalized chordwise Cp pressure tap profiles (x/c = 0.0 to 1.0) with upper/lower suction loops.",
                "• Comprehensive filtering capabilities: filter by aerospace model, flow regime, and stall status.",
                "• Extensible architecture ready for MongoDB Atlas cloud persistence and historical data archiving."
            ]
            set_shape_text(s, db_items, font_size=Pt(11.5), font_color=DARK_TEXT)
        elif s.name == "Text 7":
            sample_items = [
                "• Models: NACA 0012, SC(2)-0714, 65° Delta Wing",
                "• Facilities: Subsonic, Transonic, Supersonic",
                "• Test Runs: 156 Recorded Runs (AoA Sweeps)",
                "• Data Points: ~5,000 Calibrated Measurements",
                "• Mach Range: Mach 0.05 to 3.80",
                "• Reynolds Range: 3.0 × 10⁶ to 12.0 × 10⁶"
            ]
            set_shape_text(s, sample_items, font_size=Pt(11.5), font_color=DARK_TEXT)

    # Slide 14: API Development
    slide14 = prs.slides[13]
    for s in slide14.shapes:
        if s.has_table:
            table = s.table
            api_rows = [
                ("GET", "/api/health", "Check telemetry engine status & facility overview"),
                ("GET", "/api/facilities", "Fetch wind tunnel facility specs, RPM & pressures"),
                ("GET", "/api/models", "Retrieve aircraft model specs, polars & Cp profiles"),
                ("GET", "/api/runs", "Query and filter test runs by model, regime & stall"),
                ("GET", "/api/telemetry/live-point", "Dynamic on-the-fly aerodynamic coefficient solver"),
                ("GET", "/api/telemetry/sweep", "Generate high-resolution AoA pitch sweep arrays"),
            ]
            for r_idx, (meth, ep, purp) in enumerate(api_rows, 1):
                if r_idx < len(table.rows):
                    table.rows[r_idx].cells[0].text = meth
                    table.rows[r_idx].cells[1].text = ep
                    table.rows[r_idx].cells[2].text = purp
                    for c in table.rows[r_idx].cells:
                        for p in c.text_frame.paragraphs:
                            p.font.size = Pt(9.5)
                            p.font.color.rgb = DARK_TEXT

    # Slide 15: Module-wise Implementation Status
    slide15 = prs.slides[14]
    progress_vals = [
        ("Frontend UI & Dark HUD", "95%"),
        ("3D Canvas Flow Simulator", "90%"),
        ("Aerodynamic Analytics & Polars", "95%"),
        ("Live AoA Sweep & Audio Synth", "90%"),
        ("Backend Telemetry REST APIs", "90%"),
        ("Dataset Ingestion & Export", "85%")
    ]
    txt_idx = 0
    for s in slide15.shapes:
        if s.name == "Text 4":
            s.text_frame.text = progress_vals[0][0]
        elif s.name == "Text 7":
            s.text_frame.text = progress_vals[0][1]
        elif s.name == "Text 8":
            s.text_frame.text = progress_vals[1][0]
        elif s.name == "Text 11":
            s.text_frame.text = progress_vals[1][1]
        elif s.name == "Text 12":
            s.text_frame.text = progress_vals[2][0]
        elif s.name == "Text 15":
            s.text_frame.text = progress_vals[2][1]
        elif s.name == "Text 16":
            s.text_frame.text = progress_vals[3][0]
        elif s.name == "Text 19":
            s.text_frame.text = progress_vals[3][1]
        elif s.name == "Text 20":
            s.text_frame.text = progress_vals[4][0]
        elif s.name == "Text 23":
            s.text_frame.text = progress_vals[4][1]
        elif s.name == "Text 24":
            s.text_frame.text = progress_vals[5][0]
        elif s.name == "Text 27":
            s.text_frame.text = progress_vals[5][1]

    # Slide 16: Screenshots (Replace placeholders with 4 high-res visual previews)
    slide16 = prs.slides[15]
    # Remove text boxes Text 5, Text 7, Text 9, Text 11 and placeholders
    shapes_to_del = [s for s in slide16.shapes if s.name in ['Text 5', 'Text 7', 'Text 9', 'Text 11', 'Screenshot placeholder']]
    for shp in shapes_to_del:
        sp = shp._element
        sp.getparent().remove(sp)
    
    # Add 4 UI previews in a clean 2x2 grid
    # Top-Left: Dashboard Preview
    slide16.shapes.add_picture('ppt_assets/ui_dashboard_preview.png', Inches(0.5), Inches(1.35), Inches(4.35), Inches(1.75))
    # Top-Right: 3D Wind Tunnel Preview
    slide16.shapes.add_picture('ppt_assets/ui_windtunnel_preview.png', Inches(5.15), Inches(1.35), Inches(4.35), Inches(1.75))
    # Bottom-Left: Analytics Suite Preview
    slide16.shapes.add_picture('ppt_assets/ui_analytics_suite.png', Inches(0.5), Inches(3.25), Inches(4.35), Inches(1.75))
    # Bottom-Right: Dataset Inspector Preview
    slide16.shapes.add_picture('ppt_assets/ui_dataset_preview.png', Inches(5.15), Inches(3.25), Inches(4.35), Inches(1.75))

    # Slide 17: Challenges Faced & Solutions
    slide17 = prs.slides[16]
    for s in slide17.shapes:
        if s.has_table:
            table = s.table
            ch_data = [
                ("Real-Time Flow Simulation & Boundary Layer Separation", "Implemented an HTML5 Canvas particle vector field with Reynolds-dependent velocity damping and turbulent wake vortex shedding."),
                ("Multi-Regime Aerodynamic Coefficient Computation", "Formulated unified solvers integrating Prandtl-Glauert compressibility corrections, supercritical shock-buffet limits, and vortex lift."),
                ("High-Frequency AoA Sweep Telemetry Streaming", "Decoupled animation frame loops from React state and optimized batch updates to maintain fluid 60 FPS performance."),
                ("Cp Chordwise Pressure Distribution Integration", "Developed dual-contour upper/lower suction loop rendering with inverted aerodynamic y-axis orientation and net lift integration.")
            ]
            for r_idx, (ch, sol) in enumerate(ch_data, 1):
                if r_idx < len(table.rows):
                    table.rows[r_idx].cells[0].text = ch
                    table.rows[r_idx].cells[1].text = sol
                    for c in table.rows[r_idx].cells:
                        for p in c.text_frame.paragraphs:
                            p.font.size = Pt(9.5)
                            p.font.color.rgb = DARK_TEXT

    # Slide 18: Remaining Work & Plan
    slide18 = prs.slides[17]
    for s in slide18.shapes:
        if s.name == "Text 4":
            next_items = [
                "• Cloud Telemetry Persistence: Deploy MongoDB Atlas cluster with time-series collections for historical wind tunnel runs.",
                "• 3D Surface Pressure Contour Heatmaps: Implement WebGL / Three.js 3D isobar pressure mapping across airfoil wings.",
                "• Hardware DAQ Sensor Integration: Connect Arduino/ESP32 pitot-static pressure tubes and multi-axis load cell strain gauges.",
                "• Automated PDF Certification Export: Generate downloadable aerodynamic calibration test certificates and summary sheets.",
                "• OpenFOAM / SU2 CFD Importer: Direct ingestion of CFD mesh solutions and experimental wind tunnel data comparison."
            ]
            set_shape_text(s, next_items, font_size=Pt(12.0), font_color=DARK_TEXT)

    # Slide 19: References
    slide19 = prs.slides[18]
    for s in slide19.shapes:
        if s.has_table:
            table = s.table
            ref_data = [
                ("1", "Barlow, J. B., Rae, W. H., & Pope, A., Low-Speed Wind Tunnel Testing, 3rd Ed., John Wiley & Sons."),
                ("2", "Anderson, J. D., Fundamentals of Aerodynamics, 6th Ed., McGraw-Hill Education."),
                ("3", "Abbott, I. H., & Von Doenhoff, A. E., Theory of Wing Sections, Dover Publications."),
                ("4", "React 18, Vite 6, Chart.js, Tailwind CSS, Express.js & Web Audio API Official Documentation.")
            ]
            for r_idx, (num, ref) in enumerate(ref_data, 1):
                if r_idx < len(table.rows):
                    table.rows[r_idx].cells[0].text = num
                    table.rows[r_idx].cells[1].text = ref
                    for c in table.rows[r_idx].cells:
                        for p in c.text_frame.paragraphs:
                            p.font.size = Pt(10.0)
                            p.font.color.rgb = DARK_TEXT

    # Save to both target locations
    out_local = r'c:\tech\projects\aerodynamics\WindTunnel_Aerodynamics_Review2.pptx'
    prs.save(out_local)
    prs.save(ref_path)
    print(f"Presentation successfully built and saved to:\n1. {out_local}\n2. {ref_path}")

if __name__ == '__main__':
    build_presentation()
